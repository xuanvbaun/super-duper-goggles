import importlib.util
from pathlib import Path
import unittest
from unittest.mock import patch, MagicMock
import tempfile
import openpyxl

spec = importlib.util.spec_from_file_location('numeric_craft', Path(__file__).resolve().parents[1] / 'run.py')
craft = importlib.util.module_from_spec(spec)
spec.loader.exec_module(craft)


class NumericReviewTests(unittest.TestCase):
    def test_weight_formats_and_ambiguity(self):
        for text in ('12,78', '12 . 78', '12.78kg', '12.78 公斤'):
            self.assertEqual(craft._clean_weight(text), 12.78)
        for text in ('kg', '12 78', '1,234', '12.78 Ares', '-6'):
            self.assertIsNone(craft._clean_weight(text))

    def test_quantity_rejects_multiple_numbers(self):
        for text in ('6 9', '-6', '6/9', 'ABC6'):
            self.assertIsNone(craft._clean_qty(text))

    def test_six_nine_conflict_keeps_original(self):
        value, review = craft.quantity_decision('6', .99, [['9'], ['9'], ['9']])
        self.assertEqual(value, 6)
        self.assertTrue(review)

    def test_missing_qty_requires_three_agreements(self):
        self.assertEqual(craft.quantity_decision('', 0, [['9'], ['9'], ['9']]), (9, ''))
        self.assertIsNone(craft.quantity_decision('', 0, [['9'], [], ['9']])[0])

    def test_invalid_weight_triggers_reread(self):
        rec = {'weight': ('kg', .1), '_y': 100}
        tables = {'a': {'file': 'sample.pdf', 'pages': [[]], 'records': [rec]}}
        pdf = MagicMock()
        pdf.__len__.return_value = 1
        with patch('pypdfium2.PdfDocument', return_value=pdf), \
             patch('rapidocr_onnxruntime.RapidOCR'), \
             patch.object(craft, '_column_window', return_value=(10, 40)), \
             patch.object(craft, '_ocr_row_cell', return_value=['12.78kg']) as read:
            craft.reread_missing_weights(tables)
        read.assert_called_once()
        self.assertEqual(rec['weight'][0], '12.78')

    def test_missing_quantity_stays_blank_and_propagates(self):
        row = craft.process_doc({'meta': {}, 'records': [{'name': ('钢板', 1), 'weight': ('12.78kg', 1)}]})[0]
        self.assertIsNone(row['qty'])
        self.assertIsNone(craft.weight_per(row))
        self.assertIsNone(craft.safe_product(None, 6))

    def test_review_survives_excel_save(self):
        parent = craft.Node({'qty': 6, 'qty_review': '6/9冲突'})
        node = craft.Node({'qty': 2, 'weight': None, 'qty_review': '', 'qty_evidence': ''})
        node.parent = parent
        wb = openpyxl.Workbook()
        craft.mark_numeric_review(wb.active, 3, node)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'review.xlsx'
            wb.save(path)
            restored = openpyxl.load_workbook(path)
            for col in (7, 10, 11):
                self.assertEqual(restored.active.cell(3, col).fill.fgColor.rgb, '00FFF2CC')
                self.assertIsNotNone(restored.active.cell(3, col).comment)
            self.assertIn('待确认', restored.active.cell(3, 14).value)


if __name__ == '__main__':
    unittest.main()
