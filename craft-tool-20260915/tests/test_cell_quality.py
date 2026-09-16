import importlib.util
from pathlib import Path
import unittest
from unittest.mock import patch, MagicMock
from PIL import Image, ImageDraw
import numpy as np
import openpyxl

spec = importlib.util.spec_from_file_location('quality_craft', Path(__file__).resolve().parents[1] / 'run.py')
craft = importlib.util.module_from_spec(spec)
spec.loader.exec_module(craft)


class CellQualityTests(unittest.TestCase):
    def test_red_stamp_removed_only_from_copy(self):
        image = Image.new('RGB', (3, 1))
        image.putdata([(180, 70, 80), (10, 10, 10), (200, 200, 200)])
        cleaned = craft.prepare_ocr_image(image)
        self.assertEqual(list(cleaned.getdata()), [(255, 255, 255), (10, 10, 10), (200, 200, 200)])
        self.assertEqual(image.getpixel((0, 0)), (180, 70, 80))

    def test_spec_dimensions_remain_ordered(self):
        self.assertNotEqual(craft.spec_review_key('5x50x110'), craft.spec_review_key('50x5x110'))
        self.assertNotEqual(craft.spec_review_key('L=116.1'), craft.spec_review_key('L=1161'))

    def test_safe_spec_cleanup_applies_to_both_readings(self):
        self.assertEqual(craft.spec_review_key('K=960 80x43x5x8'),
                         craft.spec_review_key('80x43x5x8 L=960'))
        self.assertEqual(craft.spec_review_key('5×50×110'), craft.spec_review_key('5x50x110'))

    def test_grating_model_does_not_merge_with_dimension(self):
        self.assertEqual(craft.normalize_spec_spacing('G353|40|1001480x750'), 'G353|40|100 1480x750')
        self.assertEqual(craft.spec_review_key('G353|40|1001480x750'),
                         craft.spec_review_key('1480x750 G353/40/100'))
        self.assertEqual(craft.normalize_spec_spacing('G353140|100915x1400'), 'G353140|100915x1400')

    def test_noise_does_not_become_pipe_section(self):
        for s in ('/ DN20 Ø26.9x2.8. L=952', 'DN20 Ø26.9x2.8 / L=408', 'L>1096 DN40 Ø48.3x3.5'):
            expected = 'DN40 Ø48.3x3.5' if 'DN40' in s else 'DN20 Ø26.9x2.8'
            self.assertEqual(craft.base_section_spec(s, 'pipe'), expected)
        self.assertEqual(craft.base_section_spec('L=1096', 'pipe'), '')

    def test_profile_names_only_changed_for_channel_and_angle(self):
        for name, expected in [('热轧槽钢1', '槽钢'), ('热轧不等边角钢2', '角钢'), ('冷弯方形空心钢1', '冷弯方形空心钢')]:
            r = dict(name=name, spec='63x40x5 L=600', material='Q235-A', std='', cat='6')
            self.assertEqual(craft.material_of(r), expected+'63x40x5/Q235-A')
            self.assertEqual(r['name'], name)

    def test_drawing_reference_on_raw_profile_keeps_section(self):
        r = dict(name='热轧槽钢（对称制作）', spec='80x43x5x8 L=800', material='Q235-A', std='GDL161-440-28', cat='6')
        self.assertEqual(craft.material_of(r), '槽钢80x43x5x8/Q235-A')
        r['_has_children'] = True
        self.assertEqual(craft.material_of(r), 'Q235-A')

    def test_unknown_material_is_not_guessed(self):
        self.assertFalse(craft.suspicious_material('特殊耐热合金'))
        self.assertTrue(craft.suspicious_material('Q2356 X + R 245'))
        self.assertFalse(craft.suspicious_material('Q235-A'))

    def test_name_integrity_detects_loss_without_inventing_text(self):
        self.assertTrue(craft.unbalanced_name('平台13（对称制作'))
        self.assertTrue(craft.unbalanced_name('平台13）'))
        self.assertFalse(craft.unbalanced_name('平台13（对称制作）'))

    def test_missing_annotation_bracket_requires_identical_complete_peer(self):
        self.assertEqual(craft.complete_confirmed_annotation('平台13（对称制作', ['平台2（对称制作）']), '平台13（对称制作）')
        self.assertEqual(craft.complete_confirmed_annotation('平台13（对称制', ['平台2（对称制作）']), '平台13（对称制')
        self.assertEqual(craft.complete_confirmed_annotation('平台13（对称制作', []), '平台13（对称制作')

    def test_part_numbers_do_not_join_neighbor_numbers(self):
        self.assertIsNone(craft._clean_part_no('4 2'))
        self.assertEqual(craft._clean_part_no('2 . 1'), '2.1')

    def test_material_and_name_alerts_survive_excel_cells(self):
        ws = openpyxl.Workbook().active
        node = craft.Node(dict(qty=1, weight=2.4, name_review='名称异常', material_review='材质异常', review_evidence=['sample.png']))
        craft.mark_numeric_review(ws, 3, node)
        for column in (3, 9):
            self.assertIsNotNone(ws.cell(3, column).comment)
            self.assertIn('sample.png', ws.cell(3, column).comment.text)

    def test_wrong_row_product_is_repaired_without_format_change(self):
        ws = openpyxl.Workbook().active
        ws['K140'] = '=G140*J139'
        ws['K140'].number_format = 'General'
        craft.write_total_formula(ws, 140, 560)
        self.assertEqual(ws['K140'].value, '=G140*J140')
        self.assertEqual(ws['K140'].number_format, 'General')

    def test_custom_formula_not_replaced(self):
        ws = openpyxl.Workbook().active
        ws['K140'] = '=SUM(J139:J140)'
        craft.write_total_formula(ws, 140, 560)
        self.assertEqual(ws['K140'].value, '=SUM(J139:J140)')

    def test_cross_column_box_is_reread_per_cell(self):
        page = [{'text': '12.78 Ares', 'conf': .97, 'box': [[850, 80], [980, 80], [980, 100], [850, 100]]}]
        layout = {'xs': [i*100 for i in range(12)], 'ys': [0, 50, 120]}
        with patch.object(craft, '_ocr_row_cell', side_effect=[['12.78'], ['Ares']]):
            fixed = craft.split_crossing_boxes(page, layout, lambda: Image.new('RGB', (3300, 360)), MagicMock())
        self.assertEqual([it['text'] for it in fixed], ['12.78', 'Ares'])
        self.assertLess(craft.line_center(fixed[0]['box'])[0], 900)

    def test_failed_cross_column_read_preserves_original(self):
        page = [{'text': '12.78 Ares', 'conf': .97, 'box': [[850, 80], [980, 80], [980, 100], [850, 100]]}]
        layout = {'xs': [i*100 for i in range(12)], 'ys': [0, 50, 120]}
        with patch.object(craft, '_ocr_row_cell', return_value=[]):
            fixed = craft.split_crossing_boxes(page, layout, lambda: Image.new('RGB', (3300, 360)), MagicMock())
        self.assertEqual(fixed[0]['text'], page[0]['text'])

    def test_cross_column_replacement_does_not_duplicate_existing_text(self):
        page = [
            {'text': '12.78 Ares', 'conf': .97, 'box': [[850, 80], [980, 80], [980, 100], [850, 100]]},
            {'text': 'Ares', 'conf': .97, 'box': [[920, 80], [980, 80], [980, 100], [920, 100]]},
        ]
        layout = {'xs': [i*100 for i in range(12)], 'ys': [0, 50, 120]}
        with patch.object(craft, '_ocr_row_cell', side_effect=[['12.78'], ['Ares']]):
            fixed = craft.split_crossing_boxes(page, layout, lambda: Image.new('RGB', (3300, 360)), MagicMock())
        self.assertEqual([it['text'] for it in fixed], ['12.78', 'Ares'])


if __name__ == '__main__':
    unittest.main()
