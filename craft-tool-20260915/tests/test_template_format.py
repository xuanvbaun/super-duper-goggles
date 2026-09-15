# -*- coding: utf-8 -*-
"""模板格式不可改（用户规则，2026-08-13/08-17/08-18 反复强调）。

K 列公式必须沿用模板自带的 `=G*J`：2026-09-05 那轮把它改成
`=IF(OR(G="",J=""),"",G*J)`（空值返回空白），并把 G/H/J/K 的数字格式覆盖成程序自己的值，
2026-09-15 用户指出这是改模板格式，已改回。这里钉住，避免再被"优化"掉。
"""
import importlib.util
import unittest
from pathlib import Path

import openpyxl


PROGRAM = Path(__file__).resolve().parents[1] / "run.py"
spec = importlib.util.spec_from_file_location("craft_run_fmt", PROGRAM)
craft = importlib.util.module_from_spec(spec)
spec.loader.exec_module(craft)


def sheet(last_row=10):
    """模拟模板：K 列已铺好 =G*J 公式，数字格式 General。"""
    ws = openpyxl.Workbook().active
    for r in range(3, last_row + 1):
        c = ws.cell(r, 11)
        c.value = f"=G{r}*J{r}"
        c.number_format = "General"
    return ws


class TotalFormulaTests(unittest.TestCase):
    def test_template_formula_is_left_alone(self):
        ws = sheet()
        craft.write_total_formula(ws, 5, template_last_row=10)
        self.assertEqual(ws.cell(5, 11).value, "=G5*J5")
        self.assertEqual(ws.cell(5, 11).number_format, "General",
                         "模板范围内的数字格式也不能动")

    def test_row_beyond_template_gets_same_formula(self):
        ws = sheet(last_row=10)
        craft.write_total_formula(ws, 12, template_last_row=10)
        self.assertEqual(ws.cell(12, 11).value, "=G12*J12")

    def test_row_at_template_boundary_is_untouched(self):
        ws = sheet(last_row=10)
        craft.write_total_formula(ws, 10, template_last_row=10)
        self.assertEqual(ws.cell(10, 11).value, "=G10*J10")

    def test_program_never_wraps_the_formula_in_if(self):
        src = PROGRAM.read_text(encoding="utf-8")
        self.assertNotIn("=IF(OR(G", src,
                         "K 列公式必须与模板一致（=G*J），不要加 IF 包装")
